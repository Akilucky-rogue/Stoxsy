"""AMFI "TER of MF Schemes" download (manual, from AMFI's TER page) -> parsed rows. Pure: no database, no network.

The real file (checked 2026-10-30): one sheet, header on row 1, one row per scheme per DAY, the regular and the direct plan side by side (base expense ratio, brokerage, transaction cost,
statutory levies, total), the NSDL scheme code as identifier (no ISIN, AMC or AMFI numeric code), a blank row and then a disclaimer footer. Numbers are text in the workbook.
See docs/design/AMFI_TER_DESIGN.md. Values are kept exactly as published; an all-zero plan means "not offered" and is stored as missing.
"""
from __future__ import annotations

import csv
import hashlib
import io
import math
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Optional

import pandas as pd

TER_COLS = ["nsdl_code", "nsdl_code_raw", "name", "scheme_type", "scheme_category", "ter_date",
            "r_ber", "r_brokerage", "r_transaction", "r_statutory", "r_total",
            "d_ber", "d_brokerage", "d_transaction", "d_statutory", "d_total"]
PLAN_PARTS = ("ber", "brokerage", "transaction", "statutory", "total")
REJECT_MAX_SHARE = 0.02
HEADER_SCAN_ROWS = 30
EXCEL_EPOCH = datetime(1899, 12, 30)
DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d-%b-%y", "%d %b %Y", "%Y/%m/%d")


@dataclass
class Export:
    frame: pd.DataFrame
    mapping: dict
    unmatched_headers: list
    layout_hash: str
    sha256: str
    sheet: str
    n_data_rows: int
    n_parsed: int
    n_rejected: int
    rejected: list = field(default_factory=list)
    n_footer_rows: int = 0
    footer_head: str = ""


def normalise_code(raw: str) -> str:
    """The published NSDL scheme code without any whitespace and without one trailing '?'."""
    s = re.sub(r"\s+", "", str(raw))
    return s[:-1] if s.endswith("?") else s


def _norm(h) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(h).lower())


def _canonical(header) -> Optional[str]:
    """Map one source header to a column of TER_COLS (tolerant to case, spacing and punctuation), or None."""
    n = _norm(header)
    if n in ("nsdlschemecode", "nsdlcode"):
        return "nsdl_code"
    if n == "schemename":
        return "name"
    if n == "schemetype":
        return "scheme_type"
    if n == "schemecategory":
        return "scheme_category"
    if n in ("terdate", "date"):
        return "ter_date"
    plan = "r" if n.startswith("regularplan") else "d" if n.startswith("directplan") else None
    if plan is None:
        return None
    if "total" in n:
        return f"{plan}_total"
    if "brokerage" in n:
        return f"{plan}_brokerage"
    if "transaction" in n:
        return f"{plan}_transaction"
    if "statutory" in n or "levies" in n:
        return f"{plan}_statutory"
    if "base" in n or n.endswith("ber") or "expenseratio" in n:
        return f"{plan}_ber"
    return None


def _blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip()) or (isinstance(v, float) and math.isnan(v))


def parse_date(v) -> Optional[date]:
    if _blank(v):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):
        return (EXCEL_EPOCH + timedelta(days=float(v))).date() if 20000 < float(v) < 80000 else None
    s = str(v).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    try:                                                                  # '2026-10-01 00:00:00'
        return datetime.fromisoformat(s).date()
    except ValueError:
        return None


def parse_number(v) -> Optional[float]:
    if _blank(v):
        return None
    try:
        x = float(str(v).strip().replace(",", "").rstrip("%"))
    except ValueError:
        return None
    return x if math.isfinite(x) else None


def _rows_xlsx(payload: bytes):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    for ws in wb.worksheets:
        yield ws.title, [list(r) for r in ws.iter_rows(values_only=True)]


def _rows_csv(payload: bytes):
    for enc in ("utf-8-sig", "cp1252"):
        try:
            text = payload.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    yield "csv", [list(r) for r in csv.reader(io.StringIO(text))]


def _find_header(rows: list) -> Optional[int]:
    for i, r in enumerate(rows[:HEADER_SCAN_ROWS]):
        if any(_canonical(c) == "nsdl_code" for c in r if not _blank(c)):
            return i
    return None


def read_export(payload: bytes, filename: str = "") -> Export:
    """Parse an AMFI TER download (.xlsx or .csv). Raises ValueError, naming the cause, for every hard refusal."""
    sha = hashlib.sha256(payload).hexdigest()
    is_xlsx = payload[:2] == b"PK" or filename.lower().endswith((".xlsx", ".xlsm"))
    sheets = _rows_xlsx(payload) if is_xlsx else _rows_csv(payload)
    sheet, rows, h = None, None, None
    for title, rs in sheets:
        h = _find_header(rs)
        if h is not None:
            sheet, rows = title, rs
            break
    if rows is None:
        raise ValueError("no NSDL scheme code column found in any sheet: this is not an AMFI TER download")
    header = rows[h]
    mapping, unmatched, col_of = {}, [], {}
    for j, cell in enumerate(header):
        if _blank(cell):
            continue
        c = _canonical(cell)
        if c is None or c in col_of:
            unmatched.append(str(cell))
        else:
            mapping[str(cell)] = c
            col_of[c] = j
    missing = [c for c in ["nsdl_code", "name", "ter_date"] + [f"{p}_{x}" for p in "rd" for x in PLAN_PARTS] if c not in col_of]
    if missing:
        raise ValueError(f"TER column(s) missing after mapping: {', '.join(missing)}")
    layout_hash = hashlib.sha256("|".join(_norm(c) for c in header if not _blank(c)).encode()).hexdigest()

    def cell(r, c):
        j = col_of.get(c)
        return r[j] if j is not None and j < len(r) else None

    data, rejected = [], []
    n_data, end = 0, len(rows)
    for i in range(h + 1, len(rows)):
        r = rows[i]
        if all(_blank(v) for v in r):                                      # the blank row that precedes the footer
            end = i
            break
        n_data += 1
        code, name, d = cell(r, "nsdl_code"), cell(r, "name"), parse_date(cell(r, "ter_date"))
        nums = {c: parse_number(cell(r, c)) for c in col_of if c[:2] in ("r_", "d_")}
        reason = ("blank scheme code" if _blank(code) else "blank scheme name" if _blank(name) else "unreadable TER date" if d is None
                  else "unreadable or blank number in " + next(c for c, v in nums.items() if v is None) if any(v is None for v in nums.values()) else None)
        if reason:
            rejected.append({"row": i + 1, "reason": reason})
            continue
        for p in "rd":                                                      # all five values 0 = the plan is not offered
            if all(nums[f"{p}_{x}"] == 0 for x in PLAN_PARTS):
                for x in PLAN_PARTS:
                    nums[f"{p}_{x}"] = float("nan")
        data.append({"nsdl_code_raw": str(code), "nsdl_code": normalise_code(code), "name": str(name).strip(),
                     "scheme_type": None if _blank(cell(r, "scheme_type")) else str(cell(r, "scheme_type")).strip(),
                     "scheme_category": None if _blank(cell(r, "scheme_category")) else str(cell(r, "scheme_category")).strip(),
                     "ter_date": d, **nums})
    footer = rows[end:]
    footer_head = ""
    for r in footer:
        if not all(_blank(v) for v in r):
            if not _blank(cell(r, "nsdl_code")) and parse_date(cell(r, "ter_date")) is not None:
                raise ValueError("data after the footer: a row with a scheme code and a date follows the blank row that ends the data")
            if not footer_head:
                footer_head = next(str(v).strip() for v in r if not _blank(v))
    if n_data == 0:
        raise ValueError("no data rows found below the header")
    if len(rejected) > REJECT_MAX_SHARE * n_data:
        raise ValueError(f"{len(rejected)} of {n_data} data rows rejected (more than {REJECT_MAX_SHARE:.0%}): unreadable dates or numbers; first reason: {rejected[0]['reason']}")
    if not data:
        raise ValueError("no parseable rows: every TER date or number was unreadable")
    frame = pd.DataFrame(data, columns=TER_COLS)
    frame["ter_date"] = pd.to_datetime(frame["ter_date"])
    return Export(frame=frame, mapping=mapping, unmatched_headers=unmatched, layout_hash=layout_hash, sha256=sha, sheet=sheet, n_data_rows=n_data,
                  n_parsed=len(frame), n_rejected=len(rejected), rejected=rejected, n_footer_rows=len(footer), footer_head=footer_head)


TOTAL_MAX = 3.0                          # percent: a published total above this is flagged (not corrected)
SUM_TOL = 0.02                           # percent: total must equal base + brokerage + transaction + statutory within this
FLAGS = ("total_gt_3_regular", "total_gt_3_direct", "direct_not_below_regular", "sum_mismatch")
PARTIAL_SHARE = 0.70                     # a date with fewer rows than this share of the median of the previous five dates is partial
PARTIAL_LOOKBACK = 5


def flag_rows(frame: pd.DataFrame) -> pd.Series:
    """Quality flags per row (list of FLAGS names, aligned to frame.index). Values are never changed; a plan that is not offered (missing) is never flagged."""
    r_on, d_on = frame["r_total"].notna(), frame["d_total"].notna()
    r_sum = frame[["r_ber", "r_brokerage", "r_transaction", "r_statutory"]].sum(axis=1, min_count=4)
    d_sum = frame[["d_ber", "d_brokerage", "d_transaction", "d_statutory"]].sum(axis=1, min_count=4)
    masks = {"total_gt_3_regular": (r_on & (frame["r_total"] > TOTAL_MAX)),
             "total_gt_3_direct": (d_on & (frame["d_total"] > TOTAL_MAX)),
             "direct_not_below_regular": (r_on & d_on & (frame["d_total"] >= frame["r_total"])),
             "sum_mismatch": ((r_on & ((frame["r_total"] - r_sum).abs() > SUM_TOL)) | (d_on & ((frame["d_total"] - d_sum).abs() > SUM_TOL)))}
    arr = [masks[k].to_numpy() for k in FLAGS]
    return pd.Series([[k for k, a in zip(FLAGS, row) if a] for row in zip(*arr)], index=frame.index, dtype=object)


def partial_dates(frame: pd.DataFrame) -> list:
    """Dates whose row count is below 70 percent of the median of the previous five dates (a download taken before AMFI finished the day). The first five dates are never partial."""
    counts = frame.groupby(frame["ter_date"].dt.date).size().sort_index()
    out = []
    for i in range(PARTIAL_LOOKBACK, len(counts)):
        if counts.iloc[i] < PARTIAL_SHARE * counts.iloc[i - PARTIAL_LOOKBACK:i].median():
            out.append(counts.index[i])
    return out


# ---- matching TER schemes (NSDL codes) to mf_scheme_master ----------------------------------------------------------------------------

KEY_VERSION = "name_key_v1"
ACTIVE_DAYS = 90
SUGGEST_CUTOFF = 0.92
_PLAN_WORDS = re.compile(r"\b(?:direct|regular)\s*(?:plan)?\b|\bgrowth\b|\bidcw\b|\bdividend\b|\b(?:payout|re-?\s?investment|reinvestment)\b|\boption\b|\bplan\b|\bbonus\b"
                         r"|\b(?:weekly|monthly|quarterly|daily|annual|fortnightly|half[ -]?yearly)\b|\bincome distribution cum capital withdrawal\b|\bcapital withdrawal\b")
_GROWTH = re.compile(r"\bgrowth\b")
_IDCW = re.compile(r"\b(?:idcw|dividend|payout|reinvestment|income distribution)\b")


def name_key(name) -> str:
    """name_key_v1: lowercase, '&' to 'and', plan and option words removed, then letters and digits only."""
    s = str(name).lower().replace("&", " and ")
    return re.sub(r"[^a-z0-9]+", "", _PLAN_WORDS.sub(" ", s))


def option_class(name) -> Optional[str]:
    """'growth' or 'idcw' when the scheme name says exactly one of them, else None."""
    s = str(name).lower()
    g, d = bool(_GROWTH.search(s)), bool(_IDCW.search(s))
    return "growth" if g and not d else "idcw" if d and not g else None


def ter_schemes(frame: pd.DataFrame) -> pd.DataFrame:
    """One row per NSDL code: latest name and category, and which plans are offered (any non-missing total in the file)."""
    g = frame.sort_values("ter_date").groupby("nsdl_code")
    out = g.agg(name=("name", "last"), scheme_category=("scheme_category", "last"),
                r_offered=("r_total", lambda s: bool(s.notna().any())), d_offered=("d_total", lambda s: bool(s.notna().any())))
    return out.reset_index()


@dataclass
class Match:
    mapped: pd.DataFrame
    unmatched: pd.DataFrame
    suggestions: pd.DataFrame


def match_schemes(ter: pd.DataFrame, master: pd.DataFrame, today: date) -> Match:
    """Join TER schemes to master schemes by name key. Nothing is guessed: every code and every active master scheme ends up mapped or listed with a reason.
    TER reasons: no_name_key, ambiguous_key, no_plan_match. Master reasons (active schemes only): not_in_file, ambiguous_key, plan_unknown, plan_not_offered."""
    import difflib

    t = ter.copy()
    t["key"] = t["name"].map(name_key)
    m = master.copy()
    m["key"] = m["name"].map(name_key)
    m["active"] = pd.to_datetime(m["last_nav_date"]) >= pd.Timestamp(today) - pd.Timedelta(days=ACTIVE_DAYS)
    by_key = {k: g for k, g in m.groupby("key")}
    mapped, um, master_reason = [], [], {}

    def plan_for(row, r_on, d_on):
        """The TER plan that applies to a master scheme. ETFs have no plan: they take the one plan the file fills in. Otherwise its own plan when known (and offered), else the single offered plan."""
        etf = getattr(row, "is_etf", None)
        if etf is not None and not pd.isna(etf) and bool(etf) and r_on != d_on:
            return ("regular" if r_on else "direct", None)
        d = row.is_direct
        if d is not None and not pd.isna(d):
            plan, on = ("direct", d_on) if bool(d) else ("regular", r_on)
            return (plan, None) if on else (None, "plan_not_offered")
        if r_on != d_on:
            return ("regular" if r_on else "direct", None)
        return (None, "plan_unknown")

    for key, tg in t.groupby("key"):
        mg = by_key.get(key)
        if mg is None:
            um += [("ter", r.nsdl_code, None, r.name, "no_name_key") for r in tg.itertuples()]
            continue
        pairs = []
        if len(tg) == 1:
            pairs = [(tg.iloc[0], mg)]
        else:
            oc = tg["name"].map(option_class)
            if oc.notna().all() and oc.nunique() == len(tg):
                pairs = [(r, mg[mg["is_growth"] == (c == "growth")]) for (_, r), c in zip(tg.iterrows(), oc)]
            else:
                um += [("ter", r.nsdl_code, None, r.name, "ambiguous_key") for r in tg.itertuples()]
        used = set()
        for r, rows in pairs:
            n_before = len(mapped)
            for x in rows.itertuples():
                used.add(x.scheme_code)
                plan, why = plan_for(x, bool(r["r_offered"]), bool(r["d_offered"]))
                if plan:
                    mapped.append((r["nsdl_code"], int(x.scheme_code), plan))
                else:
                    master_reason[x.scheme_code] = why
            if len(mapped) == n_before:
                um.append(("ter", r["nsdl_code"], None, r["name"], "ambiguous_key" if rows.empty else "no_plan_match"))
        for x in mg.itertuples():
            if x.scheme_code not in used:
                master_reason[x.scheme_code] = "ambiguous_key"
    mapped_df = pd.DataFrame(mapped, columns=["nsdl_code", "scheme_code", "plan"])
    got = set(mapped_df["scheme_code"])
    keys_in_file = set(t["key"])
    for x in m[m["active"] & ~m["scheme_code"].isin(got)].itertuples():
        um.append(("master", None, int(x.scheme_code), x.name, master_reason.get(x.scheme_code) or ("not_in_file" if x.key not in keys_in_file else "ambiguous_key")))
    unmatched = pd.DataFrame(um, columns=["side", "nsdl_code", "scheme_code", "name", "reason"])
    mkeys = list(by_key)
    sug = []
    for r in unmatched[(unmatched["side"] == "ter") & (unmatched["reason"] == "no_name_key")].itertuples():
        for c in difflib.get_close_matches(name_key(r.name), mkeys, n=3, cutoff=SUGGEST_CUTOFF):
            sug.append((r.nsdl_code, r.name, by_key[c]["name"].iloc[0]))
    suggestions = pd.DataFrame(sug, columns=["nsdl_code", "name", "candidate_master_name"])
    return Match(mapped=mapped_df, unmatched=unmatched, suggestions=suggestions)
